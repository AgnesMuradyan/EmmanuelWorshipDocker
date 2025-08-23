from django.http import HttpResponse
from rest_framework import viewsets, status
from rest_framework.response import Response
from rest_framework.decorators import action
from .models import Album, Song, Instrument, Musician, MusicianInstrument, Singer, Plan, PlanSong
from .serializers import (
    AlbumSerializer,
    SongSerializer,
    InstrumentSerializer,
    MusicianSerializer,
    MusicianInstrumentSerializer,
    SingerSerializer,
    PlanSerializer, PlanSongSerializer
)
from django.http import HttpResponse
from django.views.decorators.http import require_GET
from unicodedata import normalize
from django.shortcuts import render, get_object_or_404
from rest_framework.views import APIView
from rest_framework import viewsets
from rest_framework.filters import SearchFilter, OrderingFilter
from .models import Song
from .serializers import SongSerializer
from rest_framework.decorators import action
from rest_framework.response import Response
from .serializers import SongSerializer, SongChoiceSerializer
from io import BytesIO
from django.http import HttpResponse
from docx import Document
from io import BytesIO
from django.http import HttpResponse
from rest_framework.decorators import action
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
# views.py
from django.http import HttpResponse
from django.views.decorators.csrf import csrf_exempt
from unicodedata import normalize

# pptx bits (already used elsewhere in your project)
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
# views.py
from django.http import JsonResponse
from django.views.decorators.csrf import ensure_csrf_cookie
ALLOWED_ORIGIN = "*"  # or set to "http://localhost:3000" for stricter dev


def index(request):
    return render(request, 'index.html')


class AlbumViewSet(viewsets.ModelViewSet):
    queryset = Album.objects.all()
    serializer_class = AlbumSerializer


class SongViewSet(viewsets.ModelViewSet):
    queryset = Song.objects.all().order_by("title")
    serializer_class = SongSerializer
    filter_backends = [SearchFilter, OrderingFilter]
    search_fields = ["title", "album__title"]
    ordering_fields = ["title", "created_at", "id"]

    @action(detail=False, methods=["get"], url_path="choices")
    def choices(self, request):
        qs = self.filter_queryset(self.get_queryset()) \
                 .order_by("title") \
                 .values("id", "title")  # ✅ returns dicts straight from DB

        page = self.paginate_queryset(qs)
        if page is not None:
            return self.get_paginated_response(list(page))
        return Response(list(qs))

    @action(detail=True, methods=['get'], url_path='view-chords', url_name='view_chords')
    def view_chords(self, request, pk=None):
        song = self.get_object()
        if not song.chords:
            return Response({'status': 'no chords available'}, status=status.HTTP_404_NOT_FOUND)
        response = HttpResponse(song.chords, content_type='application/pdf')
        response['Content-Disposition'] = f'inline; filename="{song.title}_chords.pdf"'
        return response

    @action(detail=True, methods=['get'], url_path='view-powerpoint', url_name='view_powerpoint')
    def view_powerpoint(self, request, pk=None):
        song = self.get_object()
        if not song.powerpoint:
            return Response({'status': 'no PowerPoint available'}, status=status.HTTP_404_NOT_FOUND)
        response = HttpResponse(song.powerpoint, content_type='application/vnd.openxmlformats-officedocument.presentationml.presentation')
        response['Content-Disposition'] = f'attachment; filename="{song.title}_powerpoint.pptx"'
        return response


class PlanViewSet(viewsets.ModelViewSet):
    queryset = Plan.objects.all()
    serializer_class = PlanSerializer
    filter_backends = [OrderingFilter]
    ordering_fields = ["date", "id", "day_type"]
    ordering = ["-date", "-id"]

    def get_queryset(self):
        qs = super().get_queryset()
        # existing filters:
        date_val = self.request.query_params.get("date")
        if date_val:
            qs = qs.filter(date=date_val)

        dt_csv = self.request.query_params.get("day_type")
        if dt_csv:
            wanted = [v.strip().upper() for v in dt_csv.split(",") if v.strip()]
            qs = qs.filter(day_type__in=wanted)
        return qs

    @action(detail=False, methods=["get"], url_path="choices")
    def choices(self, request):
        qs = self.filter_queryset(self.get_queryset()) \
            .order_by("-date", "-id") \
            .values("id", "date", "day_type")
        page = self.paginate_queryset(qs)
        if page is not None:
            return self.get_paginated_response(list(page))
        return Response(list(qs))

    @action(detail=True, methods=['get'], url_path='view-songs', url_name='view_songs')
    def view_songs(self, request, pk=None):
        plan = self.get_object()
        if not plan.songs.exists():
            return Response({'status': 'no songs available'}, status=status.HTTP_404_NOT_FOUND)
        songs = plan.songs.all()
        return Response({'songs': [song.song_title for song in songs]})

    @action(detail=True, methods=['get'], url_path='view-musicians', url_name='view_musicians')
    def view_musicians(self, request, pk=None):
        plan = self.get_object()
        if not plan.musicians.exists():
            return Response({'status': 'no musicians available'}, status=status.HTTP_404_NOT_FOUND)
        musicians = plan.musicians.all()
        return Response({'musicians': [f'{musician.first_name} {musician.last_name}' for musician in musicians]})

    @action(detail=True, methods=['get'], url_path='view-singers', url_name='view_singers')
    def view_singers(self, request, pk=None):
        plan = self.get_object()
        if not plan.singers.exists():
            return Response({'status': 'no singers available'}, status=status.HTTP_404_NOT_FOUND)
        singers = plan.singers.all()
        return Response({'singers': [f'{singer.first_name} {singer.last_name}' for singer in singers]})

    @action(detail=True, methods=['get'], url_path='view-lead-singers', url_name='view_lead_singers')
    def view_lead_singers(self, request, pk=None):
        plan = self.get_object()
        if not plan.lead_singers.exists():
            return Response({'status': 'no lead singers available'}, status=status.HTTP_404_NOT_FOUND)
        lead_singers = plan.lead_singers.all()
        return Response({'lead_singers': [f'{singer.first_name} {singer.last_name}' for singer in lead_singers]})

    @action(detail=True, methods=['get'], url_path='download-concatenated-powerpoint',
            url_name='download_concatenated_powerpoint')
    def download_concatenated_powerpoint(self, request, pk=None):
        plan = self.get_object()
        pptx_data = plan.create_concatenated_powerpoint()

        if not pptx_data:
            return Response({'status': 'no concatenated PowerPoint available'}, status=status.HTTP_404_NOT_FOUND)

        response = HttpResponse(pptx_data,
                                content_type='application/vnd.openxmlformats-officedocument.presentationml.presentation')
        response['Content-Disposition'] = f'attachment; filename="Plan_{plan.date}_concatenated_powerpoint.pptx"'
        return response

    @action(detail=True, methods=['get'], url_path='download-summary-docx', url_name='download_summary_docx')
    def download_summary_docx(self, request, pk=None):
        # Prefetch to avoid N+1 queries
        plan = (
            self.get_queryset()
            .prefetch_related('lead_singers', 'singers', 'songs')
            .get(pk=pk)
        )

        doc = Document()

        # Set global font to Calibri 16
        normal = doc.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(16)

        def full_names(qs):
            names = []
            for obj in qs:
                first = getattr(obj, 'first_name', '') or ''
                last = getattr(obj, 'last_name', '') or ''
                n = f"{first} {last}".strip()
                if n:
                    names.append(n)
            return names

        def comma_join(items):
            return ', '.join(items) if items else '—'

        def add_bold_line(text):
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.bold = True
            r.font.name = 'Calibri'
            r.font.size = Pt(16)
            return p

        def add_label_value(label_armenian, value_text):
            """Label (bold 16) + value (regular 16) on same line."""
            p = doc.add_paragraph()
            r_label = p.add_run(f"{label_armenian}՝ ")
            r_label.bold = True
            r_label.font.name = 'Calibri'
            r_label.font.size = Pt(16)

            r_val = p.add_run(value_text if value_text else '—')
            r_val.bold = False
            r_val.font.name = 'Calibri'
            r_val.font.size = Pt(16)
            return p

        def add_horizontal_rule():
            """Adds a stand-alone paragraph that renders a horizontal line (bottom border)."""
            p = doc.add_paragraph()
            p_el = p._element
            pPr = p_el.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '12')  # thickness (1/8 pt units)
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'auto')
            pBdr.append(bottom)
            pPr.append(pBdr)
            return p

        # --- Date line (bold, 16) ---
        try:
            # Armenian-style punctuation for date: dd.MM․yy
            date_str = plan.date.strftime('%d.%m․%y')
        except Exception:
            date_str = str(plan.date)
        add_bold_line(date_str)

        # --- Վարողներ (lead singers) ---
        lead_qs = plan.planleadsinger_set.select_related('singer').order_by('order')
        lead_names = [f"{pls.singer.first_name} {pls.singer.last_name}".strip() for pls in lead_qs]
        add_label_value('Վարողներ', ', '.join(lead_names) if lead_names else '—')

        # --- Վոկալ (singers) ---
        vocal_text = comma_join(full_names(plan.singers.all()))
        add_label_value('Վոկալ', vocal_text)

        # --- spacer then a horizontal line before "Երգեր" ---
        doc.add_paragraph("")  # blank line (not bold, 16 via Normal)
        add_horizontal_rule()

        # --- "Երգեր՝" title (bold, 16) ---
        add_bold_line("Երգեր՝")

        # Preserve custom ordering if available
        try:
            songs_qs = plan.songs.all().order_by('plansong__order', 'id')
        except Exception:
            songs_qs = plan.songs.all()

        titles = []
        for s in songs_qs:
            title = getattr(s, 'song_title', None) or getattr(s, 'title', None) or str(s)
            titles.append(title)

        if titles:
            for t in titles:
                p = doc.add_paragraph(t, style='List Number')  # numbered list items
                # Ensure numbering items are Calibri 16 (inherits from Normal, but set explicitly for safety)
                for r in p.runs:
                    r.bold = False
                    r.font.name = 'Calibri'
                    r.font.size = Pt(16)
        else:
            p = doc.add_paragraph('—')
            for r in p.runs:
                r.bold = False
                r.font.name = 'Calibri'
                r.font.size = Pt(16)

        # Serialize to bytes
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)

        resp = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        resp['Content-Disposition'] = f'attachment; filename="Plan_{plan.date}_summary.docx"'
        return resp



class InstrumentViewSet(viewsets.ModelViewSet):
    queryset = Instrument.objects.all()
    serializer_class = InstrumentSerializer


class MusicianViewSet(viewsets.ModelViewSet):
    queryset = Musician.objects.all()
    serializer_class = MusicianSerializer


class MusicianInstrumentViewSet(viewsets.ModelViewSet):
    queryset = MusicianInstrument.objects.all()
    serializer_class = MusicianInstrumentSerializer


class SingerViewSet(viewsets.ModelViewSet):
    queryset = Singer.objects.all()
    serializer_class = SingerSerializer


class PlanSongViewSet(viewsets.ModelViewSet):
    queryset = PlanSong.objects.all()
    serializer_class = PlanSongSerializer


# class PlanViewSet(viewsets.ModelViewSet):
#     queryset = Plan.objects.all()
#     serializer_class = PlanSerializer
#
# class PlanList(APIView):
#     def get(self, request):
#         plans = Plan.objects.all()
#         serializer = PlanSerializer(plans, many=True)
#         return Response(serializer.data)
#
# class PlanDetail(APIView):
#     def get(self, request, pk):
#         plan = get_object_or_404(Plan, pk=pk)
#         serializer = PlanSerializer(plan)
#         return Response(serializer.data)



# class SongList(APIView):
#     def get(self, request):
#         songs = Song.objects.all()
#         serializer = SongSerializer(songs, many=True)
#         return Response(serializer.data)
#
#
# class SongDetail(APIView):
#     def get(self, request, pk):
#         song = get_object_or_404(Song, pk=pk)
#         serializer = SongSerializer(song)
#         return Response(serializer.data)



@require_GET
def create_slide_dl(request):
    try:
        text = request.GET.get("text", "")
        build_font = request.GET.get("font", "Arial Armenian")   # font used while building
        final_font = request.GET.get("final_font", "Agg-Book1")  # forced at the end

        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
            from pptx.oxml.xmlchemy import OxmlElement
            from pptx.oxml.ns import qn
        except Exception as e:
            return HttpResponse(f"python-pptx not available: {e}", status=500, content_type="text/plain; charset=utf-8")

        # ---- Unicode -> ArmSCII-8 bytes -> Latin-1 (with «→§ and »→¦) ----
        def unicode_to_armscii8_latin1(s: str) -> str:
            s = normalize("NFC", s)
            if not hasattr(unicode_to_armscii8_latin1, "_map"):
                pairs = [
                    ('Ա','ա'),('Բ','բ'),('Գ','գ'),('Դ','դ'),('Ե','ե'),('Զ','զ'),('Է','է'),('Ը','ը'),
                    ('Թ','թ'),('Ժ','ժ'),('Ի','ի'),('Լ','լ'),('Խ','խ'),('Ծ','ծ'),('Կ','կ'),('Հ','հ'),
                    ('Ձ','ձ'),('Ղ','ղ'),('Ճ','ճ'),('Մ','մ'),('Յ','յ'),('Ն','ն'),('Շ','շ'),
                    ('Ո','ո'),('Չ','չ'),('Պ','պ'),('Ջ','ջ'),('Ռ','ռ'),('Ս','ս'),('Վ','վ'),('Տ','տ'),
                    ('Ր','ր'),('Ց','ց'),('Ւ','ւ'),('Փ','փ'),('Ք','ք'),('Օ','օ'),('Ֆ','ֆ'),
                ]
                code = 0xB2
                m = {}
                for up, lo in pairs:
                    m[ord(up)] = code; code += 1
                    m[ord(lo)] = code; code += 1
                m[ord('և')] = 0xA2
                m[ord('։')] = 0xA3
                m[ord('՝')] = 0xAA
                m[ord('֊')] = 0xAD
                m[ord('…')] = 0xAE
                m[ord('՜')] = 0xAF
                m[ord('՛')] = 0xB0
                m[ord('՞')] = 0xB1
                m[ord('՚')] = 0xFE
                unicode_to_armscii8_latin1._map = m

            out = bytearray(); mp = unicode_to_armscii8_latin1._map
            for ch in s:
                cp = ord(ch)
                if cp in mp:
                    out.append(mp[cp])
                elif cp == 0x00AB:      # «
                    out.append(0xA7)     # §
                elif cp == 0x00BB:      # »
                    out.append(0xA6)     # ¦
                elif cp <= 0xFF:
                    out.append(cp)       # pass other Latin-1
                else:
                    out.append(ord('?'))
            return out.decode("latin-1")

        # ---- helpers (fonts/theme) ----
        def _set_rfonts(el, name: str):
            rFonts = el.find(qn('a:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('a:rFonts'); el.insert(0, rFonts)
            for key in ('ascii','hAnsi','ea','cs'):
                rFonts.set(qn(f'a:{key}'), name)
                theme_attr = qn(f'a:{key}Theme')
                if theme_attr in rFonts.attrib:
                    del rFonts.attrib[theme_attr]

        def _override_theme_fonts(prs, name: str):
            try:
                theme = prs.part.theme_part._element
                elems = theme.find(qn('a:themeElements')) or None
                if elems is None: return
                scheme = elems.find(qn('a:fontScheme')) or None
                if scheme is None: return
                for tag in ('a:majorFont','a:minorFont'):
                    node = scheme.find(qn(tag)) or OxmlElement(tag)
                    if node.getparent() is None: scheme.append(node)
                    latin = node.find(qn('a:latin')) or OxmlElement('a:latin')
                    if latin.getparent() is None: node.insert(0, latin)
                    latin.set('typeface', name)
                    for sub in ('a:ea','a:cs'):
                        sub_el = node.find(qn(sub)) or OxmlElement(sub)
                        if sub_el.getparent() is None: node.append(sub_el)
                        sub_el.set('typeface', name)
            except Exception:
                pass

        # ---- parse input (2 lines per slide) ----
        raw_lines = [normalize("NFC", ln.strip())
                     for ln in text.replace("\r\n","\n").split("\n") if ln.strip()]
        legacy_lines = [unicode_to_armscii8_latin1(ln) for ln in raw_lines]
        chunks = [legacy_lines[i:i+2] for i in range(0, len(legacy_lines), 2)] or [[""]]

        prs = Presentation()
        prs.slide_width = Inches(23.00)
        prs.slide_height = Inches(12.00)

        # first: EMPTY slide 1
        s = prs.slides.add_slide(prs.slide_layouts[5])
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(0,0,0)

        # use build_font while composing
        _override_theme_fonts(prs, build_font)

        # content slides
        for chunk in chunks:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor(0,0,0)

            tb = slide.shapes.add_textbox(Inches(1.0), Inches(0.0), Inches(21.0), prs.slide_height)
            tf = tb.text_frame; tf.clear()
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
            tf.auto_size = MSO_AUTO_SIZE.NONE

            def add_line(s, first=False):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                try:
                    p.space_before = Pt(0); p.space_after = Pt(0)
                except Exception:
                    pass
                r = p.add_run()
                r.text = s
                r.font.name = build_font
                r.font.size = Pt(81)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255,255,255)
                _set_rfonts(r._r.get_or_add_rPr(), build_font)

            add_line(chunk[0], first=True)
            if len(chunk) > 1:
                add_line(chunk[1])

        # last: EMPTY slide
        s2 = prs.slides.add_slide(prs.slide_layouts[5])
        bg2 = s2.background.fill; bg2.solid(); bg2.fore_color.rgb = RGBColor(0,0,0)

        # ---- FINAL PASS: force EVERYTHING to final_font (Agg-Book1) ----
        _override_theme_fonts(prs, final_font)
        for slide in prs.slides:
            for shape in slide.shapes:
                # text frames on shapes/placeholders
                if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                    tf = shape.text_frame
                    # set default char props too
                    txBody = tf._txBody
                    lstStyle = txBody.find(qn('a:lstStyle')) or OxmlElement('a:lstStyle')
                    if lstStyle.getparent() is None: txBody.append(lstStyle)
                    defRPr = lstStyle.find(qn('a:defRPr')) or OxmlElement('a:defRPr')
                    if defRPr.getparent() is None: lstStyle.append(defRPr)
                    _set_rfonts(defRPr, final_font)
                    defRPr.set('sz', str(int(81 * 100)))
                    defRPr.set('b', '1')
                    for p in tf.paragraphs:
                        for run in p.runs:
                            run.font.name = final_font
                            run.font.size = Pt(81)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255,255,255)
                            _set_rfonts(run._r.get_or_add_rPr(), final_font)
                # tables (cells have text frames)
                if hasattr(shape, "has_table") and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            tf = cell.text_frame
                            for p in tf.paragraphs:
                                for run in p.runs:
                                    run.font.name = final_font
                                    run.font.size = Pt(81)
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255,255,255)
                                    _set_rfonts(run._r.get_or_add_rPr(), final_font)

        # ---- return file ----
        from io import BytesIO
        from time import strftime
        buf = BytesIO(); prs.save(buf); buf.seek(0)
        filename = f"slides_{strftime('%Y%m%d_%H%M%S')}.pptx"
        resp = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp

    except Exception as e:
        return HttpResponse(f"Error: {e}", status=500, content_type="text/plain; charset=utf-8")